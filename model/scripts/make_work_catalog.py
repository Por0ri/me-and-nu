# make_work_catalog — me;nu 작품 번호 목록을 만든다
#
# AniList·MAL 은 시즌·극장판을 모두 따로 센다 (진격의 거인만 해도 여러 개). 팬은 "진격의 거인" 한 덩어리로 생각하고
# 글도 대부분 시리즈 이야기라서, me;nu 는 시리즈 단위로 작품 번호를 붙인다. 시즌·극장판은 AniList 번호를 그대로 쓴다.
# - 묶는 규칙: 속편·전편 관계 + 제목 앞부분 일치 (model/algo/series.py. EXP-03 에서 MAL 로 검증한 규칙)
# - 번호: 인기 순으로 1번부터. 한 번 붙은 번호는 바꾸지 않는다 (목록이 이미 있으면 기존 번호를 두고 새 시리즈만 뒤에 붙임)
# - 한국어 제목: 위키데이터 → AniList 다른 이름(한글) 순서로 찾고, 사람이 확인한 이름(ko_titles_manual.xlsx)이 있으면 그걸 쓴다.
#   확실하지 않은 이름은 추측해서 채우지 않고 title_ko_check 에 표시한다 (입력 필요 = 빈칸 / 확인 필요 = 이름은 있지만 확인 안 됨)
#
#   python model/scripts/fetch_anilist.py links        # 관계·인기 (먼저)
#   python model/scripts/fetch_korean_titles.py        # 위키데이터 한국어 제목 (먼저)
#   python model/scripts/make_work_catalog.py
#   → model/data/menu_works.csv (작품 = 시리즈) · menu_work_entries.csv (시즌·극장판) (AniList 자료라 커밋 안 함)
#   → model/data/ko_titles_manual.xlsx: 사람이 한글 제목을 확인하는 엑셀 (인기 순. 쓰는 법은 파일 안 '사용법' 시트)

import csv
import json
import pathlib
import re
import sys
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from model.algo.series import build_series  # noqa: E402

DATA = REPO / "model" / "data"
ANILIST = DATA / "anilist"
OUT_WORKS = DATA / "menu_works.csv"
OUT_ENTRIES = DATA / "menu_work_entries.csv"
MANUAL_KO = DATA / "ko_titles_manual.xlsx"   # 사람이 한글 제목을 확인하는 엑셀 (위 설명)
NEED_INPUT, NEED_CHECK, DONE = "입력 필요", "확인 필요", "완료"
WORKLIST = [  # 엑셀 칸: (이름, 머리글, 너비)
    ("menu_work_id", "작품 번호", 9), ("rank", "인기 순위", 9), ("priority", "우선순위", 9), ("status", "상태", 11),
    ("title_ko", "한글 제목 (여기에 입력)", 40), ("search", "검색", 12), ("title_native", "일본어 원제", 30),
    ("title_english", "영어 제목", 30), ("guess_ko", "Claude 추정 (참고만)", 30), ("title_romaji", "로마자 제목", 30),
    ("first_year", "첫 해", 7), ("anilist_url", "AniList", 10), ("note", "메모", 24),
    ("original", "처음 이름 (고치지 않기)", 30)]
GUIDE = [
    "한글 제목 확인 목록 쓰는 법",
    "",
    "1. '할 일' 시트를 위에서부터 보기 (인기 순). 1순위(인기 300위 안)부터, 시간 되는 만큼만 해도 됨",
    "2. 상태가 '입력 필요'(빨강): [네이버 검색]을 눌러 정식 한글 제목을 찾아 '한글 제목' 칸에 적기 → 적으면 자동으로 완료",
    "3. 상태가 '확인 필요'(노랑): '한글 제목'이 정식 제목과 같은지 보기. 맞으면 상태를 '완료'로, 틀리면 고쳐 쓰기 → 자동으로 완료",
    "4. 한국 정식 제목이 아예 없으면: '한글 제목'을 비우고 상태를 '완료'로",
    "5. 정식 제목 기준: 국내 OTT(라프텔·넷플릭스·티빙·왓챠) 제목 > 국내 정발 만화·소설 제목. 띄어쓰기·문장부호까지 그대로",
    "6. 'Claude 추정'은 검색용 힌트. 틀린 게 많아서 그대로 쓰면 안 됨",
    "7. '처음 이름' 칸은 고치지 않기 (무엇을 고쳤는지 알아보는 칸). 정렬·필터는 해도 되고, 줄은 지우지 않기",
    "8. 다 했으면 저장(Ctrl+S)하고 Claude에게 말하기 → 작품 목록(menu_works.csv)에 반영",
]
SEASON_RELATIONS = {"SEQUEL", "PREQUEL"}
FORMAT_ORDER = {"TV": 0, "TV_SHORT": 1, "ONA": 2, "MOVIE": 3, "OVA": 4, "SPECIAL": 5, "MUSIC": 6}
HANGUL = re.compile(r"[가-힣]")


def read_csv(path):
    # 엑셀에서 그냥 저장하면 cp949 로 바뀌어서 둘 다 읽는다
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            pass
    raise SystemExit(f"{path.name} 를 읽을 수 없음 → 엑셀이면 'CSV UTF-8' 형식으로 저장")


def write_csv(path, rows, cols):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        wr = csv.DictWriter(f, fieldnames=cols)
        wr.writeheader()
        wr.writerows(rows)


def load_entries():
    entries = {}
    for name in ("media.jsonl", "media_recent.jsonl"):
        with open(ANILIST / name, encoding="utf-8") as f:
            for line in f:
                for m in json.loads(line)["media"]:
                    if not m["isAdult"]:
                        entries[m["id"]] = m
    # 처음 받은 파일(media.jsonl)엔 형식·시작 연도가 없어서 MAL 작품 표에서 MAL 번호로 채운다
    mal = {}
    with open(DATA / "mal_sample" / "anime.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            mal[int(r["anime_id"])] = r
    for m in entries.values():
        r = mal.get(m.get("idMal"))
        if "format" not in m:
            m["format"] = r["type"].upper().replace(" ", "_") if r and r["type"] else None
        if "startDate" not in m:
            m["startDate"] = {"year": int(r["start_date"][:4])} if r and r["start_date"] else None
    links = {}
    with open(ANILIST / "links.jsonl", encoding="utf-8") as f:
        for line in f:
            for m in json.loads(line)["media"]:
                links[m["id"]] = m
    return entries, links


WIKI_NOTE = re.compile(r"\s*\((?:[^)]*(?:애니메이션|애니|TV|시리즈|영화|드라마)|\d{4}년)[^)]*\)\s*$")   # "(2011년 애니메이션)"
SEASON_TAIL = re.compile(r"\s*(?:제?\s*\d+\s*기|시즌\s*\d+|\d+\s*화|파트\s*\d+|part\s*\d+)\s*$", re.I)
MOVIE_HEAD = re.compile(r"^(?:극장판|劇場版)\s+")
MOVIE_TAIL = re.compile(r"\s+(?:극장판|劇場版)$")                                               # "XXX 홀릭 극장판"
SUBTITLE = re.compile(r"\s+[-~～].*$")                                                       # " -오디널·스케일-"
LATIN_TAIL = re.compile(r"(?<=[가-힣!?])\s+[A-Z0-9][A-Z0-9 '!?.\-]*$")                        # "드래곤볼 Z"
NOT_SERIES = re.compile(r"극장판|OVA|스페셜|특별편|총집편|외전")


def clean_ko(s):
    # 위키 구분용 꼬리표와 시즌·회차 표시를 뗀다: "나루토 (애니메이션)" → "나루토", "○○ 1기" → "○○"
    s = WIKI_NOTE.sub("", s.strip())
    prev = None
    while prev != s:
        prev = s
        s = SEASON_TAIL.sub("", s).strip()
    return s


def series_ko(s):
    # 시리즈 이름용: 극장판 표시·부제까지 뗀다 ("극장판 소드 아트 온라인 -오디널·스케일-" → "소드 아트 온라인",
    # "귀멸의 칼날: 유곽 편" → "귀멸의 칼날"). 콜론 앞이 영문 2글자 이하면 그대로 ("Re: 제로부터 …"), 한글은 2글자도 자름 ("간츠: 오")
    s = SUBTITLE.sub("", MOVIE_HEAD.sub("", clean_ko(s)))
    head = s.split(":")[0].strip()
    short = head.replace(" ", "")
    if len(short) >= 3 or (len(short) == 2 and HANGUL.search(short)):
        s = head
    return clean_ko(MOVIE_TAIL.sub("", s))


def korean_candidates(m, link, wiki):
    # 시즌 하나의 한국어 이름 후보: 위키데이터(AniList 번호 → MAL 번호) → AniList 다른 이름 중 한글
    for key, val in (("anilist", m["id"]), ("mal", m.get("idMal"))):
        for hit in wiki[key].get(str(val), []) if val else []:
            yield clean_ko(hit["ko"]), "wikidata"
    for s in (link or {}).get("synonyms") or []:
        if HANGUL.search(s):
            yield clean_ko(s), "anilist"


def drop_season_tail(name, first):
    # 끝의 영문·숫자 중 첫 시즌 원제에 없는 단어는 시즌 부제라서 뗀다 ("노라가미 ARAGOTO" → "노라가미",
    # "기동전사 건담 SEED DESTINY" → "기동전사 건담 SEED"). "뱀파이어 헌터 D", "블렌드 S" 처럼 원제에 있으면 그대로
    m = LATIN_TAIL.search(name)
    if not m:
        return name
    title = f"{first['title']['romaji']} {first['title']['english'] or ''}".lower()
    known = {t.strip(".-") for t in re.findall(r"[a-z0-9.\-]+", title)}
    words = m.group().split()
    while words and words[-1].lower().strip("!?.'-") not in known:
        words.pop()
    return " ".join([name[:m.start()]] + words)


def pick_series_ko(members, entries, links, wiki, pop):
    """시리즈 한국어 이름 → (이름, 출처, 확인 필요 여부).
    위키데이터 · TV판 · 인기 시즌 이름을 먼저, 같은 시즌 안에서는 긴 이름(줄임말 말고 정식 이름)을 먼저.
    TV판이 있는 시리즈인데 극장판·OVA 이름에서 왔거나, AniList 별명에서 왔으면 확인 필요"""
    tv = lambda a: entries[a]["format"] in ("TV", "TV_SHORT", "ONA")
    has_tv = any(tv(a) for a in members)
    cands = []
    for a in members:
        for ko, src in korean_candidates(entries[a], links.get(a), wiki):
            full = series_ko(ko)
            if full:
                name = drop_season_tail(full, entries[members[0]])
                side = has_tv and (bool(NOT_SERIES.search(ko)) or not tv(a))   # TV판이 있는데 극장판·OVA 이름
                cands.append((name, src, side, a, name == full))
    if not cands:
        return "", "", True
    best = min(cands, key=lambda c: (c[2], c[1] != "wikidata", -pop[c[3]], -len(c[0])))
    # 고른 이름이 다른 확실한 이름으로 시작하면 그게 시리즈 이름 ("기동전사 건담 철혈의 오펀스" → "기동전사 건담")
    bases = [c for c in cands if c[4] and not c[2] and c[1] == "wikidata" and best[0].startswith(c[0] + " ")]
    if bases:
        best = min(bases, key=lambda c: len(c[0]))
    return best[0], best[1], best[2] or best[1] != "wikidata"


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_worklist(path):
    """확인 엑셀 → {작품 번호: 줄}. 머리글 이름으로 칸을 찾아서 정렬·칸 순서가 바뀌어도 된다"""
    rows = load_workbook(path, data_only=True)["할 일"].iter_rows(values_only=True)
    keys = {head: key for key, head, _ in WORKLIST}
    cols = [keys.get(cell_text(h)) for h in next(rows)]
    out = {}
    for row in rows:
        r = {k: cell_text(v) for k, v in zip(cols, row) if k}
        if r.get("menu_work_id", "").isdigit():
            out[int(r["menu_work_id"])] = r
    return out


def worklist_check(r):
    # 상태를 완료로 바꿨거나 한글 제목을 고쳐 쓴 줄은 확인 끝. 아니면 비어 있으면 입력 필요, 있으면 확인 필요
    name = r.get("title_ko", "")
    if r.get("status") == DONE or (name and name != r.get("original", "")):
        return ""
    return NEED_INPUT if not name else NEED_CHECK


def write_worklist(work_rows, manual, main_id):
    """확인할 작품(입력 필요 · 확인 필요)과 사람이 이미 본 작품을 인기 순으로 엑셀에 쓴다.
    한글 제목 · 상태는 menu_works.csv 와 똑같이 맞추고, Claude 추정 · 메모는 그대로 둔다"""
    rows = []
    for rank, r in enumerate(sorted(work_rows, key=lambda r: -r["popularity"]), 1):
        k = r["menu_work_id"]
        m = manual.get(k, {})
        if not m and not r["title_ko_check"]:
            continue
        query = r["title_ko"] or m.get("guess_ko") or r["title_native"] or r["title_romaji"]
        rows.append({**m, "menu_work_id": k, "rank": rank, "priority": f"{1 if rank <= 300 else 2 if rank <= 1000 else 3}순위",
                     "status": r["title_ko_check"] or DONE, "title_ko": r["title_ko"], "original": r["title_ko"],
                     "search": "https://search.naver.com/search.naver?query=" + quote(f"{query} 애니"),
                     "title_native": r["title_native"], "title_english": r["title_english"],
                     "title_romaji": r["title_romaji"], "first_year": r["first_year"],
                     "anilist_url": f"https://anilist.co/anime/{main_id[k]}"})
    rows.sort(key=lambda r: r["status"] == DONE)                # 할 일을 위로, 확인 끝난 줄은 아래로 (각각 인기 순)
    known = {r["menu_work_id"] for r in work_rows}
    rows += [m for k, m in manual.items() if k not in known]   # 없어진 번호도 지우지 않음

    wb = Workbook()
    ws = wb.active
    ws.title = "할 일"
    ws.append([head for _, head, _ in WORKLIST])
    for r in rows:
        ws.append([r.get(key, "") for key, _, _ in WORKLIST])
    col = {key: get_column_letter(i) for i, (key, _, _) in enumerate(WORKLIST, 1)}
    last, n = col[WORKLIST[-1][0]], len(rows) + 1
    for key, _, width in WORKLIST:
        ws.column_dimensions[col[key]].width = width
        head = ws[f"{col[key]}1"]
        head.font = Font(bold=True)
        head.fill = PatternFill("solid", fgColor="FFE699" if key == "title_ko" else "D9D9D9")
    for i in range(2, n + 1):                                  # 링크는 눌러서 바로 열리게
        for key, text in (("search", "네이버 검색"), ("anilist_url", "AniList")):
            c = ws[f"{col[key]}{i}"]
            if c.value:
                c.hyperlink, c.value, c.font = c.value, text, Font(color="0563C1", underline="single")
    s = col["status"]
    pick = DataValidation(type="list", formula1=f'"{NEED_INPUT},{NEED_CHECK},{DONE}"', allow_blank=True)
    ws.add_data_validation(pick)
    pick.add(f"{s}2:{s}{n}")
    fill = lambda color: PatternFill(start_color=color, end_color=color, fill_type="solid")
    ws.conditional_formatting.add(f"A2:{last}{n}", FormulaRule(formula=[f'${s}2="{DONE}"'], fill=fill("EDEDED")))
    for value, color in ((NEED_INPUT, "F8CBAD"), (NEED_CHECK, "FFF2CC")):
        ws.conditional_formatting.add(f"{s}2:{s}{n}", FormulaRule(formula=[f'{s}2="{value}"'], fill=fill(color)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last}{n}"
    guide = wb.create_sheet("사용법")
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    try:
        wb.save(MANUAL_KO)
    except PermissionError:
        print(f"  {MANUAL_KO.name} 가 엑셀에 열려 있어서 못 고침 → 엑셀을 닫고 다시 실행")
    return rows


def main():
    entries, links = load_entries()
    wiki = json.loads((DATA / "wikidata_ko.json").read_text(encoding="utf-8"))
    pairs = [(a, e["node"]["id"]) for a, l in links.items() if a in entries
             for e in l["relations"]["edges"]
             if e["relationType"] in SEASON_RELATIONS and e["node"]["type"] == "ANIME" and e["node"]["id"] in entries]
    root = build_series({a: m["title"]["romaji"] for a, m in entries.items()}, pairs)
    groups = {}
    for a in entries:
        groups.setdefault(root.get(a, a), []).append(a)

    pop = {a: (links.get(a) or {}).get("popularity") or 0 for a in entries}
    year = lambda a: (entries[a]["startDate"] or {}).get("year") or 9999
    order = lambda a: (year(a), FORMAT_ORDER.get(entries[a]["format"], 9), a)
    works = []
    for members in groups.values():
        members.sort(key=order)
        works.append({"members": members, "popularity": max(pop[a] for a in members)})

    # 번호: 이미 있는 목록의 번호는 그대로, 새 시리즈만 인기 순으로 뒤에
    old = {}
    if OUT_WORKS.exists():
        for r in read_csv(OUT_WORKS):
            for a in r["anilist_ids"].split(";"):
                old[int(a)] = int(r["menu_work_id"])
    next_id = max(old.values(), default=0) + 1
    for w in sorted(works, key=lambda w: (-w["popularity"], w["members"][0])):
        hits = sorted({old[a] for a in w["members"] if a in old})
        if len(hits) > 1:
            print(f"  주의: 예전 작품 번호 {hits} 가 한 시리즈로 합쳐짐 → {hits[0]} 사용", flush=True)
        if hits:
            w["id"] = hits[0]
        else:
            w["id"], next_id = next_id, next_id + 1

    manual = read_worklist(MANUAL_KO) if MANUAL_KO.exists() else {}
    entry_rows, work_rows, main_id = [], [], {}
    for w in sorted(works, key=lambda w: w["id"]):
        ko, src, weak = pick_series_ko(w["members"], entries, links, wiki, pop)
        check = NEED_INPUT if not ko else NEED_CHECK if weak else ""
        m = manual.get(w["id"])
        if m:                                                         # 사람이 보는 엑셀에 있으면 그 내용이 우선
            if m.get("title_ko", "") != ko:
                src = "manual" if m.get("title_ko") else ""
            ko, check = m.get("title_ko", ""), worklist_check(m)
        first = entries[max(w["members"], key=lambda a: (pop[a], -a))]   # 대표 제목: 가장 인기 있는 시즌
        main_id[w["id"]] = first["id"]
        years = [year(a) for a in w["members"] if year(a) != 9999]
        work_rows.append({
            "menu_work_id": w["id"], "title_ko": ko, "title_ko_source": src, "title_ko_check": check,
            "title_romaji": first["title"]["romaji"], "title_english": first["title"]["english"] or "",
            "title_native": first["title"]["native"] or "", "first_year": min(years, default=""),
            "last_year": max(years, default=""), "n_entries": len(w["members"]), "popularity": w["popularity"],
            "anilist_ids": ";".join(map(str, w["members"]))})
        for a in w["members"]:
            m = entries[a]
            e_ko, _ = next(korean_candidates(m, links.get(a), wiki), ("", ""))
            entry_rows.append({
                "anilist_id": a, "mal_id": m.get("idMal") or "", "menu_work_id": w["id"], "format": m["format"] or "",
                "year": year(a) if year(a) != 9999 else "", "title_ko": e_ko, "title_romaji": m["title"]["romaji"],
                "title_english": m["title"]["english"] or "", "title_native": m["title"]["native"] or "",
                "popularity": pop[a]})

    # 이름이 겹치면(리메이크·극장판이 따로 묶인 경우) 인기 1위만 그대로 두고, 나머지는 첫 해를 붙여 확인 필요로
    same = {}
    for r in work_rows:
        if r["title_ko"]:
            same.setdefault(r["title_ko"], []).append(r)
    for rows in same.values():
        for r in sorted(rows, key=lambda r: -r["popularity"])[1:]:
            r["title_ko"] += f" ({r['first_year']})" if r["first_year"] != "" else ""
            r["title_ko_check"] = NEED_CHECK

    for path, rows in ((OUT_WORKS, work_rows), (OUT_ENTRIES, entry_rows)):
        write_csv(path, rows, list(rows[0]))
    todo = write_worklist(work_rows, manual, main_id)

    top = sorted(work_rows, key=lambda r: -r["popularity"])
    count = lambda rows, c: sum(1 for r in rows if r["title_ko_check"] == c)
    print(f"작품(시리즈) {len(work_rows):,}개 · 시즌·극장판 {len(entry_rows):,}개 (속편·전편 관계 {len(pairs):,}쌍)")
    for name, rows in (("전체", work_rows), ("인기 상위 300", top[:300]), ("인기 상위 1000", top[:1000])):
        print(f"한국어 제목 {name}: 문제없음 {count(rows, ''):,} · {NEED_CHECK} {count(rows, NEED_CHECK):,} · "
              f"{NEED_INPUT} {count(rows, NEED_INPUT):,}")
    print(f"→ {OUT_WORKS}\n→ {OUT_ENTRIES}\n→ {MANUAL_KO} "
          f"(할 일 {sum(1 for r in todo if r.get('status') != DONE):,}줄 · 확인 끝 {sum(1 for r in todo if r.get('status') == DONE):,}줄)")


if __name__ == "__main__":
    main()
